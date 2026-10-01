#!/usr/bin/env python3
"""slide-studio B2 render engine — spike

構造化スライドデータ + 画風テーマ → 同一のレイアウト計算(px座標)から
  (A) HTML を描いて PNG に(Playwright/headless Chromium)
  (B) python-pptx で「編集可能な」PPTX(実テキストボックス)
を両出力する。両バックエンドが同じ box を消費するので見た目が一致する(parity)。

K1(フォント一致): デッキで使う文字だけに subset した「同一フォントバイト」を
  PNG = HTML @font-face(base64) / PPTX = 埋め込み(embeddedFont)
の両方に使う → Hiragino 非搭載のPCでも PNG と同じ書体で表示される。

これが B2(編集可能PowerPoint)の背骨。gpt-image は使わない(課金ゼロ)。
"""
from __future__ import annotations
import argparse, base64, json, pathlib, sys
import yaml
import fonts   # engine/fonts.py
import icons   # engine/icons.py

ROOT = pathlib.Path(__file__).resolve().parents[1]

PX_TO_EMU = 9525
def emu(px): return int(round(px * PX_TO_EMU))
def pt(px):  return px * 0.75


# ============================================================
# レイアウト(両バックエンド共有)
# ============================================================
def _style(theme, role, color_key):
    t = theme["typography"][role]
    return {"family": t["family"], "size": t["size"], "weight": t.get("weight", 400),
            "line": t.get("line", 1.3), "tracking": t.get("tracking", 0),
            "upper": t.get("upper", False), "color": theme["palette"][color_key]}

def _mix(h1, h2, t):
    a = [int(h1.lstrip("#")[i:i+2], 16) for i in (0, 2, 4)]
    b = [int(h2.lstrip("#")[i:i+2], 16) for i in (0, 2, 4)]
    return "#%02x%02x%02x" % tuple(round(a[i] + (b[i] - a[i]) * t) for i in range(3))

def card_colors(theme):
    p = theme["palette"]; bg = theme["canvas"]["bg"]
    return {"fill": _mix(bg, p["accent"], 0.045), "border": p.get("rule", _mix(bg, p["heading"], 0.10)),
            "accent": p["accent"], "accent_wash": _mix(bg, p["accent"], 0.14),
            "heading": p["heading"], "body": p["body"], "muted": p["muted"], "bg": bg,
            "shadow": "0 1px 2px rgba(20,24,30,.05),0 14px 32px -12px rgba(20,24,30,.13)",
            "up": p.get("good", "#2f8f5b"), "down": p.get("bad", "#c0472f")}

def _delta(cc, d):
    """delta dict → (符号色, 記号) 。up=緑/down=赤(意味色)。"""
    up = str((d or {}).get("dir", "up")) != "down"
    return (cc["up"] if up else cc["down"]), ("▲" if up else "▼")

def _txt(theme, box, text, role, ck, **extra):
    b = {"type": "text", "box": box, "text": str(text), "style": _style(theme, role, ck)}
    b.update(extra); return b

import math as _math

def _est_lines(text, font_px, width_px, cap=3):
    """見出し等の折り返し行数を推定(全角=1.0/半角=0.55 で近似)。"""
    units = sum(0.55 if ch.isascii() else 1.0 for ch in str(text))
    per_line = max(1.0, width_px / max(1.0, font_px))
    return max(1, min(cap, _math.ceil(units / per_line)))

def _header(slide_els, theme, x, W, y0, heading_cap=3):
    blocks = []; y = y0
    blocks.append({"type": "rule", "box": (x, y, 44, 4), "color": theme["palette"]["accent"]}); y += 20  # 短いアクセント棒
    if slide_els.get("kicker"):
        blocks.append(_txt(theme, (x, y, W, 28), slide_els["kicker"]["text"], "kicker", "kicker")); y += 40
    if slide_els.get("heading"):
        htext = slide_els["heading"]["text"]
        hsize = _style(theme, "heading", "heading")["size"]
        lines = _est_lines(htext, hsize, W, cap=heading_cap)      # 行数を推定して高さを可変(2行見出しの重なり防止)
        hh = int(lines * hsize * 1.30) + 6
        blocks.append(_txt(theme, (x, y, W, hh), htext, "heading", "heading")); y += hh + 14
    return blocks, y

def layout(slide, theme):
    m = theme["margin"]; cw = theme["canvas"]["w"]; ch = theme["canvas"]["h"]
    x = m["x"]; W = cw - 2 * m["x"]; y0 = m["y"]; P = theme["palette"]
    els = {e["kind"]: e for e in slide.get("elements", []) if isinstance(e, dict) and e.get("kind")}
    lay = slide.get("layout", "title-bullets"); blocks = []

    if lay in ("cover", "closing"):
        if lay == "cover":
            blocks.append({"type": "deco", "box": (0, 0, cw, ch), "kind": "glow", "color": P["accent"]})
        blocks.append({"type": "rule", "box": (x, 238, 96, 4), "color": P["accent"]}); y = 258
        if els.get("kicker"): blocks.append(_txt(theme, (x, y, W, 34), els["kicker"]["text"], "kicker", "kicker")); y += 50
        if els.get("heading"): blocks.append(_txt(theme, (x, y, W, 150), els["heading"]["text"], "heading", "heading")); y += 156
        if els.get("subhead"): blocks.append(_txt(theme, (x, y + 8, W, 66), els["subhead"]["text"], "subhead", "muted"))
    elif lay == "section":
        blocks.append({"type": "rule", "box": (x, 300, 72, 5), "color": P["accent"]}); y = 326
        if els.get("kicker"): blocks.append(_txt(theme, (x, y, W, 30), els["kicker"]["text"], "kicker", "kicker")); y += 44
        if els.get("heading"): blocks.append(_txt(theme, (x, y, W, 150), els["heading"]["text"], "heading", "heading"))
    elif lay == "statement":
        st = _style(theme, "heading", "heading"); st["size"] = int(st["size"] * 1.02); st["align"] = "center"
        if els.get("kicker"): blocks.append(_txt(theme, (x, 244, W, 30), els["kicker"]["text"], "kicker", "kicker", align="center"))
        t = (els.get("heading") or els.get("statement") or {"text": ""})["text"]
        blocks.append({"type": "text", "box": (x, 290, W, 190), "text": str(t), "style": st})
    elif lay == "title-bullets":
        hb, y = _header(els, theme, x, W, y0); blocks += hb
        if els.get("bullets"):
            st = _style(theme, "body", "body"); st["marker_color"] = theme["bullet"]["marker_color"]; st["marker"] = theme["bullet"].get("marker", "square")
            blocks.append({"type": "bullets", "box": (x, y, W, ch - y - y0), "items": els["bullets"]["items"], "style": st, "row_h": 66})
    elif lay == "big-number":
        blocks.append({"type": "deco", "box": (0, 0, cw, ch), "kind": "glow", "color": P["accent"]})
        hb, y = _header(els, theme, x, W, y0); blocks += hb
        s = els.get("stat") or {}
        blocks.append({"type": "stat", "box": (x, y + 24, W, ch - y - y0 - 24), "value": str(s.get("value", "")),
                       "unit": str(s.get("unit", "")), "caption": str(s.get("caption", "")), "delta": s.get("delta"), "theme": theme})
    elif lay == "icon-grid":
        hb, y = _header(els, theme, x, W, y0); blocks += hb
        items = (els.get("features") or {}).get("items", [])[:4]
        blocks.append({"type": "iconcards", "box": (x, y + 8, W, ch - y - y0 - 8), "cards": items, "theme": theme})
    elif lay == "comparison":
        hb, y = _header(els, theme, x, W, y0); blocks += hb
        tb = els.get("table") or {}
        blocks.append({"type": "table", "box": (x, y + 10, W, ch - y - y0 - 10),
                       "columns": tb.get("columns", []), "rows": tb.get("rows", []), "theme": theme})
    elif lay == "process":
        hb, y = _header(els, theme, x, W, y0); blocks += hb
        st = (els.get("steps") or {}).get("items", [])[:5]
        blocks.append({"type": "steps", "box": (x, y + 24, W, ch - y - y0 - 44), "steps": st, "theme": theme})
    elif lay == "timeline":
        hb, y = _header(els, theme, x, W, y0); blocks += hb
        ms = (els.get("timeline") or {}).get("items", [])[:6]
        blocks.append({"type": "timeline", "box": (x, y + 40, W, ch - y - y0 - 60), "items": ms, "theme": theme})
    elif lay == "matrix":
        hb, y = _header(els, theme, x, W, y0); blocks += hb
        m = els.get("matrix") or {}
        blocks.append({"type": "matrix", "box": (x, y + 6, W, ch - y - y0 - 6),
                       "axis_x": m.get("axis_x", {}), "axis_y": m.get("axis_y", {}), "items": m.get("items", [])[:6], "theme": theme})
    elif lay == "funnel":
        hb, y = _header(els, theme, x, W, y0); blocks += hb
        f = els.get("funnel") or {}
        blocks.append({"type": "funnel", "box": (x, y + 14, W, ch - y - y0 - 30), "stages": f.get("stages", [])[:5], "theme": theme})
    elif lay == "image-lead":
        hb, y = _header(els, theme, x, W, y0); blocks += hb
        v = els.get("visual") or {}
        blocks.append({"type": "visual", "box": (x, y + 8, W, ch - y - y0 - 8), "icon": v.get("icon"), "points": (v.get("points") or [])[:5], "theme": theme})
    elif lay == "kpi-row":
        hb, y = _header(els, theme, x, W, y0); blocks += hb
        blocks.append({"type": "kpi", "box": (x, y + 10, W, 212), "tiles": (els.get("kpi") or {}).get("tiles", [])[:4], "theme": theme})
    elif lay == "two-column":
        hb, y = _header(els, theme, x, W, y0); blocks += hb
        c = els.get("columns") or {}
        blocks.append({"type": "columns", "box": (x, y + 4, W, ch - y - y0), "left": c.get("left", {}), "right": c.get("right", {}), "theme": theme})
    elif lay == "quote":
        q = els.get("quote") or {}
        blocks.append({"type": "quote", "box": (x, 176, W, 368), "quote": str(q.get("text", "")), "name": str(q.get("name", "")), "role": str(q.get("role", "")), "theme": theme})
    elif lay == "chart":
        hb, y = _header(els, theme, x, W, y0); blocks += hb
        ins = els.get("insight", {}).get("text", "") if isinstance(els.get("insight"), dict) else ""
        spec = (els.get("chart") or {}).get("chart") or {}   # chart要素は {kind:"chart", chart:{kind,categories,series}}
        blocks.append({"type": "chart", "box": (x, y + 8, W, ch - y - y0 - (66 if ins else 8)), "spec": spec, "theme": theme})
        if ins: blocks.append(_txt(theme, (x, ch - y0 - 58, W, 52), ins, "body", "muted"))
    else:
        raise ValueError(f"unknown layout: {lay}")
    return blocks


def composite_cells(b, theme):
    """複合ブロック(columns/iconcards/steps/timeline/table/kpi)の各テキストセルを
    {box,(x,y,w,h) canvas座標, text, role, weight, family, align} で返す。
    render_pptx が各セルを配置する幾何と同一式なので、テンプレ描画と一致する。
    ALC-A(gpt-image→編集可PPTX)がこの box を使えば、複合セル本文も決定論的に
    native 化できる(=編集可能テキスト網羅ゲートを機械的に満たせる)。既存挙動は不変。"""
    t = b.get("type"); out = []
    def C(box, text, role, weight, align=None):
        text = str(text or "").strip()
        if text:
            out.append({"box": tuple(round(v) for v in box), "text": text,
                        "role": role, "weight": weight, "family": "sans", "align": align})
    if t == "columns":
        x, y, w, h = b["box"]; hh = min(h, 430); gap = 18; cw2 = (w - gap) / 2
        for si, side in enumerate((b.get("left") or {}, b.get("right") or {})):
            cx = x + si * (cw2 + gap); yy = y + 22
            if side.get("title"):
                C((cx + 24, yy, cw2 - 48, 34), side["title"], "heading", 700); yy += 46
            for p in (side.get("points") or []):
                C((cx + 40, yy, cw2 - 64, 40), p, "body", 400); yy += 42
    elif t == "iconcards":
        x, y, w, h = b["box"]; hh = min(h, 420); gap = 16; cds = b.get("cards") or []
        n = max(1, len(cds)); tw = (w - gap * (n - 1)) / n
        for i, c in enumerate(cds):
            tx = x + i * (tw + gap); ty = y + 90 if c.get("icon") else y + 26
            C((tx + 24, ty, tw - 48, 60), c.get("title", ""), "heading", 700)
            C((tx + 24, ty + 42, tw - 48, hh - (ty - y) - 42), c.get("body", ""), "body", 400)
    elif t == "steps":
        x, y, w, h = b["box"]; sts = b.get("steps") or []; n = max(1, len(sts)); cw2 = w / n
        for i, st in enumerate(sts):
            sx = x + i * cw2
            C((sx, y + 74, cw2 - 16, 44), st.get("title", ""), "heading", 700)
            C((sx, y + 118, cw2 - 16, 90), st.get("desc", ""), "body", 400)
    elif t == "timeline":
        x, y, w, h = b["box"]; ms = b.get("items") or []; n = max(1, len(ms)); cw2 = w / n
        for i, m in enumerate(ms):
            C((x + i * cw2, y, cw2, 26), m.get("date", ""), "kicker", 600, align="center")
            C((x + i * cw2, y + 48, cw2, 60), m.get("label", ""), "heading", 700, align="center")
    elif t == "kpi":
        x, y, w, h = b["box"]; tiles = b.get("tiles") or []; n = max(1, len(tiles)); gap = 16
        tw = (w - gap * (n - 1)) / n
        for i, tl in enumerate(tiles):
            tx = x + i * (tw + gap); vy = y + 80 if tl.get("icon") else y + 22
            C((tx + 22, vy, tw - 44, 64), tl.get("value", ""), "heading", 800)
            C((tx + 22, vy + 74, tw - 44, 60), tl.get("label", ""), "body", 400)
    elif t == "table":
        x, y, w, h = b["box"]; cols = b.get("columns", []); rows = b.get("rows", [])
        ncol = len(cols); nrow = len(rows) + 1
        if ncol and rows:
            rh = min(h, 58 + len(rows) * 52) / max(1, nrow); cwid = w / ncol
            for ci in range(ncol):
                C((x + ci * cwid, y, cwid, rh), cols[ci], "heading", 700,
                  align=("center" if ci > 0 else None))
            for ri, r in enumerate(rows, 1):
                ry = y + ri * rh
                C((x, ry, cwid, rh), r.get("label", ""), "heading", 700)
                cells = r.get("cells", [])
                for ci in range(1, ncol):
                    cv = cells[ci - 1] if ci - 1 < len(cells) else ""
                    C((x + ci * cwid, ry, cwid, rh), cv, "body", 400, align="center")
    return out


# ============================================================
# フォント資産(subset)= PNG/PPTX共有
# ============================================================
def _iter_block_faces(blocks):
    """(family_kw, weight, text) を列挙(subset対象の全文字を集める)"""
    for b in blocks:
        t = b["type"]
        if t == "text":
            yield b["style"]["family"], b["style"]["weight"], b["text"]
        elif t == "bullets":
            yield b["style"]["family"], b["style"]["weight"], "".join(b["items"])
        elif t == "stat":
            yield "sans", 700, f"{b['value']}{b['unit']}"
            yield "sans", 400, b["caption"]
            d = b.get("delta")
            if isinstance(d, dict) and d.get("value"):
                yield "sans", 700, str(d["value"])   # delta は 600(bold)描画なので 700 で subset
        elif t == "kpi":
            for tl in b["tiles"]:
                yield "sans", 700, str(tl.get("value", ""))
                yield "sans", 400, str(tl.get("label", ""))
                d = tl.get("delta")
                if isinstance(d, dict) and d.get("value"):
                    yield "sans", 700, str(d["value"])
        elif t == "columns":
            for side in ("left", "right"):
                c = b.get(side) or {}
                yield "sans", 700, str(c.get("title", ""))
                yield "sans", 400, "".join(str(p) for p in (c.get("points") or []))
        elif t == "quote":
            yield "serif", 700, b["quote"]
            yield "sans", 700, b["name"]; yield "sans", 400, b["role"]
        elif t == "chart":
            sp = b.get("spec") or {}
            yield "sans", 400, "".join(str(c) for c in sp.get("categories", []))
            for s in sp.get("series", []):
                yield "sans", 700, str(s.get("name", ""))
        elif t == "iconcards":
            for c in b["cards"]:
                yield "sans", 700, str(c.get("title", ""))
                yield "sans", 400, str(c.get("body", ""))
        elif t == "table":
            for c in b.get("columns", []):
                yield "sans", 700, str(c)
            for r in b.get("rows", []):
                yield "sans", 700, str(r.get("label", ""))
                for cell in r.get("cells", []):
                    yield "sans", 400, str(cell)
        elif t == "steps":
            for st in b["steps"]:
                yield "sans", 700, str(st.get("title", ""))
                yield "sans", 400, str(st.get("desc", ""))
        elif t == "timeline":
            for m in b["items"]:
                yield "sans", 700, str(m.get("label", ""))
                yield "sans", 400, str(m.get("date", ""))
        elif t == "matrix":
            for k in ("axis_x", "axis_y"):
                a = b.get(k) or {}
                yield "sans", 400, str(a.get("low", "")) + str(a.get("high", ""))
            for it in b.get("items", []):
                yield "sans", 700, str(it.get("label", ""))
        elif t == "funnel":
            for st in b.get("stages", []):
                yield "sans", 700, str(st.get("label", "")) + str(st.get("value", ""))
        elif t == "visual":
            for pnt in b.get("points", []):
                yield "sans", 400, str(pnt)

def build_font_assets(deck, theme):
    """使う (family,bold) を集めて subset。戻り: {(family,bold): bytes}, テキスト集合"""
    used = {}  # (family,bold) -> set(chars)
    for sl in deck["slides"]:
        for fam_kw, weight, text in _iter_block_faces(layout(sl, theme)):
            fam = fonts.family_of(fam_kw); bold = weight >= 600
            used.setdefault((fam, bold), set()).update(text)
    return {key: fonts.subset_bytes(key[0], key[1], "".join(sorted(chars)))
            for key, chars in used.items()}

def font_face_css(subsets):
    out = []
    for (fam, bold), data in subsets.items():
        b64 = base64.b64encode(data).decode()
        out.append(f"@font-face{{font-family:'{fam}';font-weight:{700 if bold else 400};font-style:normal;"
                   f"src:url(data:font/ttf;base64,{b64}) format('truetype');}}")
    return "".join(out)


# ============================================================
# (A) HTML -> PNG
# ============================================================
def _css_font(style):
    fam = fonts.family_of(style["family"])
    css = (f"font-family:'{fam}',{fonts.HTML_FALLBACK[fam]};font-size:{style['size']}px;"
           f"font-weight:{style['weight']};line-height:{style['line']};color:{style['color']};")
    if style.get("tracking"): css += f"letter-spacing:{style['tracking']}em;"
    if style.get("upper"): css += "text-transform:uppercase;"
    return css

def _esc(s): return str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

def _ff(kw):
    fam = fonts.family_of(kw); return f"'{fam}',{fonts.HTML_FALLBACK[fam]}"

def _nice_max(v):
    import math
    if v <= 0: return 1.0
    base = 10 ** math.floor(math.log10(v)); f = v / base
    for c in (1, 1.2, 1.5, 2, 2.5, 3, 4, 5, 6, 8, 10):
        if f <= c: return c * base
    return 10 * base

def _fmt(v):
    try:
        v = float(v)
        return f"{v:,.0f}" if v == int(v) else f"{v:,.1f}"
    except (TypeError, ValueError):
        return str(v)

def _f01(v):
    try:
        return max(0.0, min(1.0, float(v)))
    except (TypeError, ValueError):
        return 0.5

def _fnum(v):
    try:
        return float(str(v).replace(",", "").replace("%", "").strip())
    except (TypeError, ValueError):
        return None

def chart_palette(theme):
    p = theme["palette"]; a = p["accent"]
    return [a, _mix(a, "#ffffff", 0.4), p.get("gold", _mix(a, "#000000", 0.35)), _mix(a, "#ffffff", 0.68), _mix(a, "#000000", 0.5)]

def svg_chart(spec, box, theme):
    """インラインSVGでチャート描画。kind: column/bar/line/pie/doughnut。数値はspec由来のみ。"""
    import math
    spec = spec or {}; kind = spec.get("kind", "column")
    cats = spec.get("categories") or []
    series = [s for s in (spec.get("series") or []) if isinstance(s, dict)]
    x, y, w, h = box; fam = _ff("sans"); P = theme["palette"]   # SVGはフォールバック付き(subset漏れの空白化防止)
    mut, body = P["muted"], P["body"]; pal = chart_palette(theme)
    col = lambda i, s=None: (s or {}).get("color") or pal[i % len(pal)]
    S = [f'<svg width="{w}" height="{h}" viewBox="0 0 {w} {h}" font-family="{fam}">']

    if kind in ("pie", "doughnut"):
        data = (series[0].get("data") if series else []) or []
        total = sum(data) or 1
        cx, cy, R = w * 0.30, h / 2, min(w * 0.26, h * 0.44); r0 = R * 0.58 if kind == "doughnut" else 0
        ang = -math.pi / 2
        for i, v in enumerate(data):
            a2 = ang + 2 * math.pi * (v / total); lg = 1 if (a2 - ang) > math.pi else 0
            x1, y1 = cx + R * math.cos(ang), cy + R * math.sin(ang)
            x2, y2 = cx + R * math.cos(a2), cy + R * math.sin(a2)
            if r0:
                ix1, iy1 = cx + r0 * math.cos(a2), cy + r0 * math.sin(a2)
                ix2, iy2 = cx + r0 * math.cos(ang), cy + r0 * math.sin(ang)
                d = f"M{x1:.1f} {y1:.1f}A{R:.0f} {R:.0f} 0 {lg} 1 {x2:.1f} {y2:.1f}L{ix1:.1f} {iy1:.1f}A{r0:.0f} {r0:.0f} 0 {lg} 0 {ix2:.1f} {iy2:.1f}Z"
            else:
                d = f"M{cx:.1f} {cy:.1f}L{x1:.1f} {y1:.1f}A{R:.0f} {R:.0f} 0 {lg} 1 {x2:.1f} {y2:.1f}Z"
            S.append(f'<path d="{d}" fill="{col(i)}"/>'); ang = a2
        for i, c in enumerate(cats):
            yy = cy - (len(cats) - 1) * 15 + i * 30
            S.append(f'<rect x="{w*0.60:.0f}" y="{yy-13:.0f}" width="15" height="15" rx="3" fill="{col(i)}"/>')
            val = _fmt(data[i]) if i < len(data) else ""
            S.append(f'<text x="{w*0.60+24:.0f}" y="{yy:.0f}" font-size="19" fill="{body}">{_esc(c)}　{val}</text>')
        S.append('</svg>'); return "".join(S)

    # column / bar(横) / line
    PADL, PADB, PADT, PADR = 66, 42, 18, 14
    pw, ph = w - PADL - PADR, h - PADB - PADT
    allv = [v for s in series for v in (s.get("data") or [])] or [0]
    vmax = _nice_max(max(allv) * 1.06) if max(allv) > 0 else 1
    S.append('<defs>')
    for si, s in enumerate(series):
        c = col(si, s)
        S.append(f'<linearGradient id="bg{si}" x1="0" y1="0" x2="0" y2="1"><stop offset="0%" stop-color="{c}"/>'
                 f'<stop offset="100%" stop-color="{_mix(c, "#ffffff", 0.22)}"/></linearGradient>')
    S.append('</defs>')
    for k in range(5):
        t = vmax * k / 4; gy = PADT + ph * (1 - k / 4)
        S.append(f'<line x1="{PADL}" y1="{gy:.1f}" x2="{w-PADR}" y2="{gy:.1f}" stroke="{mut}" stroke-opacity=".14"/>')
        S.append(f'<text x="{PADL-8}" y="{gy+5:.1f}" text-anchor="end" font-size="16" fill="{mut}">{_fmt(t)}</text>')
    n = max(1, len(cats)); ns = max(1, len(series))
    if kind == "line":
        for si, s in enumerate(series):
            pts = [(PADL + pw * (ci + 0.5) / n, PADT + ph * (1 - (s.get("data") or [0])[ci] / vmax))
                   for ci in range(min(n, len(s.get("data") or [])))]
            if pts:
                path = " ".join(f'{"M" if i==0 else "L"}{px:.1f} {py:.1f}' for i, (px, py) in enumerate(pts))
                S.append(f'<path d="{path}" fill="none" stroke="{col(si,s)}" stroke-width="3"/>')
                for px, py in pts: S.append(f'<circle cx="{px:.1f}" cy="{py:.1f}" r="4.5" fill="{col(si,s)}"/>')
    else:
        gw = pw / n; bw = gw * 0.62 / ns
        for ci in range(n):
            for si, s in enumerate(series):
                data = s.get("data") or []; v = data[ci] if ci < len(data) else 0
                bh = ph * (v / vmax); bx = PADL + ci * gw + gw * 0.19 + si * bw; by = PADT + ph - bh
                S.append(f'<rect x="{bx:.1f}" y="{by:.1f}" width="{bw*0.9:.1f}" height="{max(0,bh):.1f}" rx="5" fill="url(#bg{si})"/>')
                if ns == 1:
                    S.append(f'<text x="{bx+bw*0.45:.1f}" y="{by-9:.1f}" text-anchor="middle" font-size="17" font-weight="700" fill="{col(si,s)}">{_fmt(v)}</text>')
    for ci, c in enumerate(cats):
        S.append(f'<text x="{PADL+ci*pw/n+pw/n/2:.1f}" y="{h-16:.0f}" text-anchor="middle" font-size="18" fill="{mut}">{_esc(c)}</text>')
    if ns > 1:
        lx = PADL
        for si, s in enumerate(series):
            nm = _esc(s.get("name", "")); S.append(f'<rect x="{lx:.0f}" y="{PADT-2:.0f}" width="12" height="12" rx="2" fill="{col(si,s)}"/>')
            S.append(f'<text x="{lx+18:.0f}" y="{PADT+9:.0f}" font-size="15" fill="{body}">{nm}</text>'); lx += 30 + len(nm) * 15
    S.append('</svg>'); return "".join(S)

def html_for_slide(blocks, theme, font_css, bg_data=None, scrim=None):
    cw = theme["canvas"]["w"]; ch = theme["canvas"]["h"]; bg = theme["canvas"]["bg"]
    # AI背景(gpt-image)を敷く場合は背景画像 + 可読性スクリム。テキスト/数値は下のブロックで正確に描く。
    cv_bg = f"background:{bg} center/cover no-repeat;"
    scrim_div = ""
    if bg_data:
        cv_bg = f"background:{bg} url('{bg_data}') center/cover no-repeat;"
        if scrim:
            scrim_div = f'<div style="position:absolute;inset:0;background:{scrim};"></div>'
    # -webkit-font-smoothing:antialiased は文字を細く描画し PPTX(PowerPoint)側と太さがズレるため付けない
    parts = [f'<div class="cv" style="position:relative;width:{cw}px;height:{ch}px;{cv_bg}'
             'overflow:hidden;">', scrim_div]
    for b in blocks:
        x, y, w, h = b["box"]; pos = f"position:absolute;left:{x}px;top:{y}px;width:{w}px;"
        t = b["type"]
        if t == "text":
            al = "text-align:center;" if b["style"].get("align") == "center" else ""
            # max-height+overflow: 長い見出し等が予約枠を超えて下のコンポーネントに重なるのを防ぐ(枠内でクリップ)
            parts.append(f'<div data-fit="1" style="{pos}max-height:{h}px;overflow:hidden;{al}{_css_font(b["style"])}">{_esc(b["text"])}</div>')
        elif t == "rule":
            parts.append(f'<div style="{pos}height:{h}px;background:{b["color"]};border-radius:2px;"></div>')
        elif t == "bullets":
            st = b["style"]; rows = [f'<div style="{pos}top:{y}px;">']
            mk = "border-radius:2px;" if st.get("marker") == "square" else "border-radius:50%;"
            for it in b["items"]:
                rows.append(f'<div style="display:flex;gap:16px;align-items:flex-start;height:{b["row_h"]}px;">'
                            f'<span style="width:12px;height:12px;margin-top:{int(st["size"]*0.42)}px;flex:none;'
                            f'background:{st["marker_color"]};{mk}"></span>'
                            f'<span style="{_css_font(st)}">{_esc(it)}</span></div>')
            rows.append('</div>'); parts.append("".join(rows))
        elif t == "stat":
            cc = card_colors(theme); d = b.get("delta"); dh = ""
            if isinstance(d, dict) and d.get("value"):
                dcol, sym = _delta(cc, d)
                dh = (f'<div style="margin-top:18px;align-self:flex-start;display:inline-flex;align-items:center;gap:6px;'
                      f'font:600 19px {_ff("sans")};color:{dcol};background:{_mix(cc["bg"],dcol,0.13)};padding:6px 15px;border-radius:100px">{sym} {_esc(d.get("value",""))}</div>')
            parts.append(f'<div style="{pos}height:{h}px;display:flex;flex-direction:column;justify-content:center">'
                         f'<div style="display:flex;align-items:baseline;gap:12px"><span style="font:800 140px {_ff("sans")};color:{cc["accent"]};line-height:.9;letter-spacing:-.03em">{_esc(b["value"])}</span>'
                         f'<span style="font:700 44px {_ff("sans")};color:{cc["accent"]}">{_esc(b["unit"])}</span></div>'
                         f'<div style="margin-top:16px;font:400 25px {_ff("sans")};color:{cc["muted"]};max-width:760px">{_esc(b["caption"])}</div>{dh}</div>')
        elif t == "kpi":
            cc = card_colors(theme); cards = []
            for tl in b["tiles"]:
                nm = tl.get("icon"); ic = ""
                if nm and icons.has(nm):
                    ic = (f'<div style="width:40px;height:40px;border-radius:11px;background:{cc["accent_wash"]};'
                          f'display:flex;align-items:center;justify-content:center;margin-bottom:16px">{icons.svg_inline(nm, cc["accent"], 22)}</div>')
                d = tl.get("delta"); dh = ""
                if isinstance(d, dict) and d.get("value"):
                    dcol, sym = _delta(cc, d)
                    dh = f'<div style="margin-top:9px;font:600 15px {_ff("sans")};color:{dcol}">{sym} {_esc(d.get("value",""))}</div>'
                vlen = len(str(tl.get("value", ""))); vsize = 46 if vlen <= 6 else 38 if vlen <= 9 else 30  # 長い値は自動縮小(折返し防止)
                cards.append(f'<div style="flex:1;background:{cc["fill"]};border:1px solid {cc["border"]};border-radius:16px;padding:22px 20px;overflow:hidden;box-shadow:{cc["shadow"]}">'
                             f'{ic}<div style="font:800 {vsize}px {_ff("sans")};color:{cc["accent"]};line-height:1.05;word-break:break-all">{_esc(tl.get("value",""))}</div>'
                             f'<div style="margin-top:9px;font:400 15px {_ff("sans")};color:{cc["muted"]};line-height:1.4;display:-webkit-box;-webkit-line-clamp:2;-webkit-box-orient:vertical;overflow:hidden">{_esc(tl.get("label",""))}</div>{dh}</div>')
            parts.append(f'<div style="{pos}display:flex;gap:16px;height:{h}px">{"".join(cards)}</div>')
        elif t == "columns":
            cc = card_colors(theme)
            def _colhtml(c):
                title = f'<div style="font:700 18px {_ff("sans")};color:{cc["heading"]};margin-bottom:4px">{_esc(c.get("title",""))}</div>' if c.get("title") else ""
                pts = "".join(f'<div style="display:flex;gap:10px;align-items:flex-start;margin-top:11px;font:400 17px {_ff("sans")};color:{cc["body"]};line-height:1.5">'
                              f'<span style="width:7px;height:7px;border-radius:2px;background:{cc["accent"]};flex:none;margin-top:9px"></span><span>{_esc(p)}</span></div>' for p in (c.get("points") or []))
                return f'<div style="flex:1;background:{cc["fill"]};border:1px solid {cc["border"]};border-radius:16px;padding:24px 26px;overflow:hidden;box-shadow:{cc["shadow"]}">{title}{pts}</div>'
            parts.append(f'<div style="{pos}display:flex;gap:18px;align-items:stretch">{_colhtml(b["left"])}{_colhtml(b["right"])}</div>')  # 内容高さ(空白枠防止)
        elif t == "quote":
            cc = card_colors(theme); attr = ""
            if b["name"] or b["role"]:
                role = ("　" + _esc(b["role"])) if b["role"] else ""
                attr = f'<div style="margin-top:26px;font:400 17px {_ff("sans")};color:{cc["muted"]}"><span style="font-weight:700;color:{cc["heading"]}">{_esc(b["name"])}</span>{role}</div>'
            parts.append(f'<div style="{pos}height:{h}px;overflow:hidden;display:flex;flex-direction:column;justify-content:center">'
                         f'<div style="font:800 96px {_ff("serif")};color:{cc["accent"]};line-height:.5;height:40px">“</div>'
                         f'<div style="margin-top:6px;font:700 31px {_ff("serif")};color:{cc["heading"]};line-height:1.5">{_esc(b["quote"])}</div>{attr}</div>')
        elif t == "chart":
            parts.append(f'<div style="{pos}">{svg_chart(b["spec"], b["box"], theme)}</div>')
        elif t == "deco":
            cw2, ch2, col = b["box"][2], b["box"][3], b["color"]
            gid = f"g{len(parts)}"
            if b.get("kind") == "glow":
                parts.append(f'<svg style="position:absolute;left:0;top:0;pointer-events:none" width="{cw2}" height="{ch2}">'
                             f'<defs><radialGradient id="{gid}" cx="50%" cy="50%" r="50%">'
                             f'<stop offset="0%" stop-color="{col}" stop-opacity="0.16"/><stop offset="100%" stop-color="{col}" stop-opacity="0"/></radialGradient></defs>'
                             f'<circle cx="{cw2-60}" cy="-20" r="360" fill="url(#{gid})"/>'
                             f'<circle cx="-40" cy="{ch2+40}" r="240" fill="url(#{gid})"/></svg>')
        elif t == "iconcards":
            cc = card_colors(theme); cards = []
            for c in b["cards"]:
                nm = c.get("icon"); ic = ""
                if nm and icons.has(nm):
                    ic = (f'<div style="width:46px;height:46px;border-radius:13px;background:{cc["accent_wash"]};'
                          f'display:flex;align-items:center;justify-content:center;margin-bottom:16px">{icons.svg_inline(nm, cc["accent"], 25)}</div>')
                cards.append(f'<div style="flex:1;background:{cc["fill"]};border:1px solid {cc["border"]};border-radius:16px;padding:26px 24px;overflow:hidden;box-shadow:{cc["shadow"]}">'
                             f'{ic}<div style="font:700 19px {_ff("sans")};color:{cc["heading"]};line-height:1.35">{_esc(c.get("title",""))}</div>'
                             f'<div style="margin-top:10px;font:400 15px {_ff("sans")};color:{cc["body"]};line-height:1.55">{_esc(c.get("body",""))}</div></div>')
            parts.append(f'<div style="{pos}display:flex;gap:16px;align-items:stretch">{"".join(cards)}</div>')  # 内容高さ(空白枠防止)
        elif t == "table":
            cc = card_colors(theme); cols = b.get("columns", []); rows = b.get("rows", [])
            def _cell(v, first=False):
                s = str(v)
                if s in ("○", "◯", "✓", "✔"): return f'<span style="color:{cc["up"]};font:700 22px {_ff("sans")}">✓</span>'
                if s in ("×", "✗", "✕"): return f'<span style="color:{cc["down"]};font:400 20px {_ff("sans")}">✗</span>'
                if s in ("△", "–", "-", "ー"): return f'<span style="color:{cc["muted"]};font:400 20px {_ff("sans")}">△</span>'
                c = cc["heading"] if first else cc["body"]; wt = 700 if first else 400
                return f'<span style="color:{c};font-weight:{wt}">{_esc(s)}</span>'
            head = "".join(f'<th style="padding:15px 18px;text-align:{"left" if i==0 else "center"};font:700 16px {_ff("sans")};color:{cc["heading"]}">{_esc(c)}</th>' for i, c in enumerate(cols))
            body = ""
            for r in rows:
                tds = [f'<td style="padding:14px 18px;text-align:left;border-top:1px solid {cc["border"]};font:400 17px {_ff("sans")}">{_cell(r.get("label",""), True)}</td>']
                for cv in r.get("cells", []):
                    tds.append(f'<td style="padding:14px 18px;text-align:center;border-top:1px solid {cc["border"]};font:400 17px {_ff("sans")}">{_cell(cv)}</td>')
                body += "<tr>" + "".join(tds) + "</tr>"
            parts.append(f'<div style="{pos}border:1px solid {cc["border"]};border-radius:16px;overflow:hidden;box-shadow:{cc["shadow"]}">'
                         f'<table style="width:{w}px;border-collapse:collapse;background:{cc["fill"]}"><thead style="background:{cc["accent_wash"]}"><tr>{head}</tr></thead><tbody>{body}</tbody></table></div>')
        elif t == "steps":
            cc = card_colors(theme); cells = []
            for i, st in enumerate(b["steps"]):
                cells.append(f'<div style="flex:1;padding:0 8px;overflow:hidden">'
                             f'<div style="font:800 40px {_ff("sans")};color:{cc["accent"]};line-height:1">{i+1:02d}</div>'
                             f'<div style="width:36px;height:3px;background:{cc["accent"]};border-radius:2px;margin:12px 0 15px"></div>'
                             f'<div style="font:700 18px {_ff("sans")};color:{cc["heading"]};line-height:1.35">{_esc(st.get("title",""))}</div>'
                             f'<div style="margin-top:9px;font:400 15px {_ff("sans")};color:{cc["body"]};line-height:1.5">{_esc(st.get("desc",""))}</div></div>')
            parts.append(f'<div style="{pos}display:flex;gap:12px;align-items:flex-start">{"".join(cells)}</div>')
        elif t == "timeline":
            cc = card_colors(theme); ms = b["items"]; n = max(1, len(ms))
            line = f'<div style="position:absolute;left:{w/n/2:.0f}px;right:{w/n/2:.0f}px;top:34px;height:2px;background:{cc["border"]}"></div>'
            dots = ""
            for m in ms:
                dots += (f'<div style="flex:1;text-align:center;overflow:hidden">'
                         f'<div style="font:600 15px {_ff("sans")};color:{cc["accent"]};margin-bottom:14px">{_esc(m.get("date",""))}</div>'
                         f'<div style="width:15px;height:15px;border-radius:50%;background:{cc["accent"]};margin:0 auto;box-shadow:0 0 0 4px {cc["bg"]},0 0 0 6px {cc["accent"]}"></div>'
                         f'<div style="margin-top:18px;font:700 16px {_ff("sans")};color:{cc["heading"]};line-height:1.4">{_esc(m.get("label",""))}</div></div>')
            parts.append(f'<div style="{pos}position:relative;display:flex;align-items:flex-start">{line}{dots}</div>')
        elif t == "matrix":
            cc = card_colors(theme); ax = b.get("axis_x") or {}; ay = b.get("axis_y") or {}; its = b.get("items", [])
            fam = _ff("sans"); ps = min(h - 34, 470); px0 = 138; py0 = 4   # 左マージンをY軸ラベル幅ぶん確保(見切れ防止)
            sv = [f'<svg width="{w}" height="{h}" font-family="{fam}">',
                  f'<rect x="{px0}" y="{py0}" width="{ps}" height="{ps}" rx="12" fill="{cc["fill"]}" stroke="{cc["border"]}"/>',
                  f'<line x1="{px0+ps/2:.0f}" y1="{py0}" x2="{px0+ps/2:.0f}" y2="{py0+ps}" stroke="{cc["border"]}"/>',
                  f'<line x1="{px0}" y1="{py0+ps/2:.0f}" x2="{px0+ps}" y2="{py0+ps/2:.0f}" stroke="{cc["border"]}"/>']
            for it in its:
                ix = px0 + _f01(it.get("x")) * ps; iy = py0 + (1 - _f01(it.get("y"))) * ps
                sv.append(f'<circle cx="{ix:.0f}" cy="{iy:.0f}" r="9" fill="{cc["accent"]}"/>')
                sv.append(f'<text x="{ix+15:.0f}" y="{iy+6:.0f}" font-size="16" font-weight="700" fill="{cc["heading"]}">{_esc(it.get("label",""))}</text>')
            sv.append(f'<text x="{px0}" y="{py0+ps+26:.0f}" font-size="14" fill="{cc["muted"]}">{_esc(ax.get("low",""))}</text>')
            sv.append(f'<text x="{px0+ps}" y="{py0+ps+26:.0f}" text-anchor="end" font-size="14" fill="{cc["muted"]}">{_esc(ax.get("high",""))}</text>')
            sv.append(f'<text x="8" y="{py0+ps:.0f}" font-size="13" fill="{cc["muted"]}">{_esc(ay.get("low",""))}</text>')
            sv.append(f'<text x="8" y="{py0+16:.0f}" font-size="13" fill="{cc["muted"]}">{_esc(ay.get("high",""))}</text></svg>')
            parts.append(f'<div style="{pos}">{"".join(sv)}</div>')
        elif t == "funnel":
            cc = card_colors(theme); stg = b.get("stages", []); n = max(1, len(stg)); fam = _ff("sans")
            pal = chart_palette(theme); sh = (h - 8) / n; topw = w * 0.72; cx = w / 2; minw = w * 0.14
            _fvals = [_fnum(st.get("value")) for st in stg]; _mx = max([v for v in _fvals if v] or [0])
            def _wof(i):   # 幅=値に比例(値が無ければ緩やかテーパー)。脱落を視覚化
                if i < 0 or i >= n: return minw
                v = _fvals[i]
                if v is None or _mx <= 0: return max(minw, topw * (1 - 0.55 * (i / max(1, n - 1))))
                return max(minw, topw * v / _mx)
            sv = [f'<svg width="{w}" height="{h}" font-family="{fam}">']
            for i, st in enumerate(stg):
                w0 = _wof(i); w1 = _wof(i + 1) if i + 1 < n else max(minw, w0 * 0.85)
                y0 = i * sh + 4; y1 = (i + 1) * sh
                sv.append(f'<path d="M{cx-w0/2:.0f} {y0:.0f} L{cx+w0/2:.0f} {y0:.0f} L{cx+w1/2:.0f} {y1-8:.0f} L{cx-w1/2:.0f} {y1-8:.0f} Z" fill="{pal[i%len(pal)]}"/>')
                val = ("　" + _fmt(st.get("value"))) if str(st.get("value", "")).strip() else ""
                sv.append(f'<text x="{cx:.0f}" y="{(y0+y1)/2:.0f}" text-anchor="middle" font-size="19" font-weight="700" fill="#ffffff">{_esc(st.get("label",""))}{val}</text>')
            sv.append('</svg>')
            parts.append(f'<div style="{pos}">{"".join(sv)}</div>')
        elif t == "visual":
            cc = card_colors(theme); nm = b.get("icon"); pts = b.get("points", [])
            ico = icons.svg_inline(nm, cc["accent"], 128) if (nm and icons.has(nm)) else ""
            panel = f'<div style="flex:0 0 42%;background:{cc["accent_wash"]};border-radius:20px;display:flex;align-items:center;justify-content:center">{ico}</div>'
            plist = "".join(f'<div style="display:flex;gap:12px;align-items:flex-start;margin-top:15px;font:400 19px {_ff("sans")};color:{cc["body"]};line-height:1.5">'
                            f'<span style="width:8px;height:8px;border-radius:2px;background:{cc["accent"]};flex:none;margin-top:10px"></span><span>{_esc(pnt)}</span></div>' for pnt in pts)
            parts.append(f'<div style="{pos}display:flex;gap:26px;height:{min(h,440)}px">'
                         f'<div style="flex:1;display:flex;flex-direction:column;justify-content:center">{plist}</div>{panel}</div>')
    parts.append("</div>")
    return ('<!doctype html><meta charset="utf-8"><style>*{margin:0;padding:0;box-sizing:border-box}'
            f'html,body{{margin:0}}{font_css}</style>' + "".join(parts))

# data-fit 要素を、枠(max-height)に収まるまでフォントを段階縮小(下限12px)。文字切れをブラウザ実測で根絶。
# 枠に収まるまでフォントを縮小(文字切れ根絶)。ただし許容は **比例** にする:
#   CJK太字の実グリフ高は line-height(1.2/1.16) の行ボックスより ~10% 大きく、1行の見出しでも
#   scrollHeight > clientHeight(例 68 > 62)になる。旧実装の絶対許容(+1px)ではこの比例10%を
#   吸収できず、比率が縮小しても変わらないため見出しを 52px→19px まで暴走縮小していた
#   (全スライドのタイトルが約2.7倍縮む製品全体の欠陥。実測で特定)。
#   *1.18 は単一行の実測比(sans/serif とも 1.05〜1.10)の上・真の溢れ(1行増=比1.5+)の下に置く。
AUTOFIT_JS = """() => {
  document.querySelectorAll('[data-fit]').forEach(el => {
    let fs = parseFloat(getComputedStyle(el).fontSize), g = 0;
    while (el.scrollHeight > el.clientHeight * 1.18 + 1 && fs > 12 && g < 60) { fs -= 1; el.style.fontSize = fs + 'px'; g++; }
  });
}"""

def render_pngs(deck, theme, outdir, font_css):
    from playwright.sync_api import sync_playwright
    cw, ch = theme["canvas"]["w"], theme["canvas"]["h"]; paths = []
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(viewport={"width": cw, "height": ch}, device_scale_factor=2)
        for sl in deck["slides"]:
            page.set_content(html_for_slide(layout(sl, theme), theme, font_css), wait_until="load")
            # @font-face を明示ロードして完了を待つ(遅延ロードで screenshot が先行する対策)
            page.evaluate("async () => { await Promise.all([...document.fonts].map(f => f.load())); await document.fonts.ready; }")
            page.evaluate(AUTOFIT_JS)   # 枠に収まるまでフォント縮小(文字切れ根絶)
            out = outdir / f"{sl['id']}.png"; page.locator(".cv").screenshot(path=str(out)); paths.append(out)
        browser.close()
    return paths


# ============================================================
# (B) 編集可能 PPTX(フォント埋め込み)
# ============================================================
def render_pptx(deck, theme, out_path, subsets):
    from pptx import Presentation
    from pptx.util import Emu, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.oxml.ns import qn
    def rgb(h): return RGBColor.from_string(h.lstrip("#").upper())

    def set_cjk_font(run, tf):
        """日本語は East Asian スロット(a:ea)で描画される。a:latin だけでは効かないので ea/cs も設定。"""
        rPr = run._r.get_or_add_rPr()
        anchor = rPr.find(qn("a:latin"))
        for tag in ("a:ea", "a:cs"):
            el = rPr.find(qn(tag))
            if el is None:
                el = rPr.makeelement(qn(tag), {})
                if anchor is not None:
                    anchor.addnext(el)
                else:
                    rPr.append(el)
            el.set("typeface", tf)
            anchor = el

    cw, ch = theme["canvas"]["w"], theme["canvas"]["h"]
    prs = Presentation(); prs.slide_width = Emu(emu(cw)); prs.slide_height = Emu(emu(ch))
    blank = prs.slide_layouts[6]

    def add_text(slide, box, text, style):
        x, y, w, h = box
        tb = slide.shapes.add_textbox(Emu(emu(x)), Emu(emu(y)), Emu(emu(w)), Emu(emu(h)))
        tb.shadow.inherit = False
        tf = tb.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]; p.line_spacing = style["line"]
        if style.get("align") == "center": p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = text; f = r.font
        f.size = Pt(pt(style["size"])); f.bold = style["weight"] >= 600
        f.name = fonts.family_of(style["family"]); f.color.rgb = rgb(style["color"])
        set_cjk_font(r, fonts.family_of(style["family"]))   # 日本語(a:ea)にも同じ書体

    def add_rich(slide, box, runs, line=1.05):
        x, y, w, h = box
        tb = slide.shapes.add_textbox(Emu(emu(x)), Emu(emu(y)), Emu(emu(w)), Emu(emu(h)))
        tb.shadow.inherit = False; tf = tb.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]; p.line_spacing = line
        for rn in runs:
            r = p.add_run(); r.text = str(rn["text"]); f = r.font
            f.size = Pt(pt(rn["size"])); f.bold = rn.get("weight", 400) >= 600
            fam = fonts.family_of(rn.get("family", "sans")); f.name = fam; f.color.rgb = rgb(rn["color"])
            set_cjk_font(r, fam)

    def txt1(slide, box, text, size, color, weight=400, family="sans", line=1.35, align=None):
        st = {"family": family, "size": size, "weight": weight, "line": line, "color": color}
        if align: st["align"] = align
        add_text(slide, box, str(text), st)

    def style_cell(cell, color, size, bold, fillcolor, center=False):
        cell.fill.solid(); cell.fill.fore_color.rgb = rgb(fillcolor)
        cell.margin_left = cell.margin_right = Emu(emu(16)); cell.margin_top = cell.margin_bottom = Emu(emu(10))
        p = cell.text_frame.paragraphs[0]
        if center: p.alignment = PP_ALIGN.CENTER
        for r in p.runs:
            f = r.font; f.size = Pt(pt(size)); f.bold = bold
            f.name = fonts.family_of("sans"); f.color.rgb = rgb(color); set_cjk_font(r, fonts.family_of("sans"))

    def card(slide, x, y, w, h, cc):
        c = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(emu(x)), Emu(emu(y)), Emu(emu(w)), Emu(emu(h)))
        c.shadow.inherit = False
        try: c.adjustments[0] = 0.055
        except Exception: pass
        c.fill.solid(); c.fill.fore_color.rgb = rgb(cc["fill"]); c.line.color.rgb = rgb(cc["border"]); c.line.width = Pt(1)
        return c

    def _cjk_font(font, fam):
        rPr = getattr(font, "_rPr", None)
        if rPr is None:
            rPr = getattr(font, "_element", None)
        if rPr is None:
            return
        for tag in ("a:ea", "a:cs"):
            el = rPr.find(qn(tag))
            if el is None:
                el = rPr.makeelement(qn(tag), {}); rPr.append(el)
            el.set("typeface", fam)

    def add_native_chart(slide, box, spec):
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
        x, y, w, h = box; spec = spec or {}
        cats = [str(c) for c in (spec.get("categories") or [])]
        series = [ss for ss in (spec.get("series") or []) if isinstance(ss, dict) and ss.get("data")]
        if not cats or not series:
            return  # データ無しは描かない(捏造しない)
        TYPE = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED,
                "line": XL_CHART_TYPE.LINE_MARKERS, "pie": XL_CHART_TYPE.PIE, "doughnut": XL_CHART_TYPE.DOUGHNUT}
        cd = CategoryChartData(); cd.categories = cats
        for ss in series:
            data = [(float(v) if isinstance(v, (int, float)) else 0.0) for v in ss.get("data", [])]
            cd.add_series(str(ss.get("name", "")), tuple((data + [0.0] * len(cats))[:len(cats)]))
        chart = slide.shapes.add_chart(TYPE.get(spec.get("kind", "column"), XL_CHART_TYPE.COLUMN_CLUSTERED),
                                       Emu(emu(x)), Emu(emu(y)), Emu(emu(w)), Emu(emu(h)), cd).chart
        chart.has_title = False; fam = fonts.family_of("sans"); pal = chart_palette(theme)
        chart.has_legend = len(series) > 1
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM; chart.legend.include_in_layout = False
            chart.legend.font.size = Pt(10); _cjk_font(chart.legend.font, fam)
        for i, ss in enumerate(chart.series):
            try:
                ss.format.fill.solid(); ss.format.fill.fore_color.rgb = rgb(pal[i % len(pal)])
            except Exception: pass
        try:
            plot = chart.plots[0]; plot.has_data_labels = True
            plot.data_labels.font.size = Pt(10); _cjk_font(plot.data_labels.font, fam)
            if spec.get("kind") not in ("pie", "doughnut"):
                plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
        except Exception: pass
        try:
            tl = chart.category_axis.tick_labels; tl.font.size = Pt(11); _cjk_font(tl.font, fam)
        except Exception: pass

    # アイコン(kpi/icon-grid)を色付きPNGにラスタライズ(1ブラウザで一括)
    accent = theme["palette"]["accent"]; icon_specs = set()
    for _sl in deck["slides"]:
        for _b in layout(_sl, theme):
            if _b["type"] == "kpi":
                for _tl in _b["tiles"]:
                    if _tl.get("icon") and icons.has(_tl["icon"]): icon_specs.add((_tl["icon"], accent, 22))
            elif _b["type"] == "iconcards":
                for _c in _b["cards"]:
                    if _c.get("icon") and icons.has(_c["icon"]): icon_specs.add((_c["icon"], accent, 26))
            elif _b["type"] == "visual":
                if _b.get("icon") and icons.has(_b["icon"]): icon_specs.add((_b["icon"], accent, 128))
    icon_png = icons.rasterize(icon_specs) if icon_specs else {}

    def place_icon(slide, name, x, y, boxsz, isz, cc):
        sq = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(emu(x)), Emu(emu(y)), Emu(emu(boxsz)), Emu(emu(boxsz)))
        sq.shadow.inherit = False; sq.line.fill.background()
        try: sq.adjustments[0] = 0.28
        except Exception: pass
        sq.fill.solid(); sq.fill.fore_color.rgb = rgb(cc["accent_wash"])
        png = icon_png.get((name, cc["accent"], isz))
        if png:
            off = (boxsz - isz) / 2
            slide.shapes.add_picture(str(png), Emu(emu(x + off)), Emu(emu(y + off)), Emu(emu(isz)), Emu(emu(isz)))

    for sl in deck["slides"]:
        s = prs.slides.add_slide(blank)
        bgr = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bgr.shadow.inherit = False; bgr.line.fill.background()
        bgr.fill.solid(); bgr.fill.fore_color.rgb = rgb(theme["canvas"]["bg"])
        for b in layout(sl, theme):
            if b["type"] == "text":
                add_text(s, b["box"], b["text"], b["style"])
            elif b["type"] == "rule":
                x, y, w, h = b["box"]
                sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(emu(x)), Emu(emu(y)), Emu(emu(w)), Emu(emu(h)))
                sh.shadow.inherit = False; sh.line.fill.background(); sh.fill.solid(); sh.fill.fore_color.rgb = rgb(b["color"])
            elif b["type"] == "bullets":
                x, y, w, h = b["box"]; st = b["style"]; rh = b["row_h"]
                for i, it in enumerate(b["items"]):
                    ry = y + i * rh
                    shp = MSO_SHAPE.RECTANGLE if st.get("marker") == "square" else MSO_SHAPE.OVAL
                    mk = s.shapes.add_shape(shp, Emu(emu(x)), Emu(emu(ry + st["size"] * 0.42)), Emu(emu(12)), Emu(emu(12)))
                    mk.shadow.inherit = False; mk.line.fill.background(); mk.fill.solid(); mk.fill.fore_color.rgb = rgb(st["marker_color"])
                    add_text(s, (x + 28, ry, w - 28, rh), it, st)
            elif b["type"] == "stat":
                x, y, w, h = b["box"]; cc = card_colors(theme)
                add_rich(s, (x, y + 6, w, 150), [{"text": b["value"], "size": 128, "color": cc["accent"], "weight": 800},
                                                 {"text": " " + b["unit"], "size": 42, "color": cc["accent"], "weight": 700}])
                txt1(s, (x, y + 176, w, 44), b["caption"], 25, cc["muted"], 400)
                d = b.get("delta")
                if isinstance(d, dict) and d.get("value"):
                    dcol, sym = _delta(cc, d)
                    txt1(s, (x, y + 226, w, 34), sym + " " + str(d.get("value", "")), 20, dcol, 600)
            elif b["type"] == "kpi":
                x, y, w, h = b["box"]; cc = card_colors(theme); tiles = b["tiles"]; n = max(1, len(tiles)); gap = 16
                tw = (w - gap * (n - 1)) / n
                for i, tl in enumerate(tiles):
                    tx = x + i * (tw + gap); card(s, tx, y, tw, h, cc); vy = y + 22
                    if tl.get("icon") and icons.has(tl["icon"]):
                        place_icon(s, tl["icon"], tx + 22, y + 22, 40, 22, cc); vy = y + 80
                    txt1(s, (tx + 22, vy, tw - 44, 64), tl.get("value", ""), 46, cc["accent"], 800, line=1.05)
                    txt1(s, (tx + 22, vy + 74, tw - 44, 60), tl.get("label", ""), 15, cc["muted"], 400)
                    d = tl.get("delta")
                    if isinstance(d, dict) and d.get("value"):
                        dcol, sym = _delta(cc, d)
                        txt1(s, (tx + 22, y + h - 36, tw - 44, 28), sym + " " + str(d.get("value", "")), 15, dcol, 600)
            elif b["type"] == "iconcards":
                x, y, w, h = b["box"]; cc = card_colors(theme); cds = b["cards"]; n = max(1, len(cds)); gap = 16
                hh = min(h, 420); tw = (w - gap * (n - 1)) / n
                for i, c in enumerate(cds):
                    tx = x + i * (tw + gap); card(s, tx, y, tw, hh, cc); ty = y + 26
                    if c.get("icon") and icons.has(c["icon"]):
                        place_icon(s, c["icon"], tx + 24, y + 24, 46, 26, cc); ty = y + 90
                    txt1(s, (tx + 24, ty, tw - 48, 60), c.get("title", ""), 19, cc["heading"], 700, line=1.35)
                    txt1(s, (tx + 24, ty + 42, tw - 48, hh - (ty - y) - 42), c.get("body", ""), 15, cc["body"], 400, line=1.5)
            elif b["type"] == "columns":
                x, y, w, h = b["box"]; cc = card_colors(theme); hh = min(h, 430); gap = 18; cw2 = (w - gap) / 2
                for si, side in enumerate((b["left"], b["right"])):
                    cx = x + si * (cw2 + gap); card(s, cx, y, cw2, hh, cc); yy = y + 22
                    if side.get("title"):
                        txt1(s, (cx + 24, yy, cw2 - 48, 34), side["title"], 18, cc["heading"], 700); yy += 46
                    for pt_ in (side.get("points") or []):
                        mk = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(emu(cx + 24)), Emu(emu(yy + 9)), Emu(emu(8)), Emu(emu(8)))
                        mk.shadow.inherit = False; mk.line.fill.background(); mk.fill.solid(); mk.fill.fore_color.rgb = rgb(cc["accent"])
                        txt1(s, (cx + 40, yy, cw2 - 64, 40), pt_, 17, cc["body"], 400, line=1.5); yy += 42
            elif b["type"] == "quote":
                x, y, w, h = b["box"]; cc = card_colors(theme)
                txt1(s, (x, y + 4, w, 80), "“", 90, cc["accent"], 800, family="serif", line=0.9)
                txt1(s, (x, y + 92, w, 180), b["quote"], 31, cc["heading"], 700, family="serif", line=1.5)
                if b["name"] or b["role"]:
                    attr = str(b["name"]) + (("　" + str(b["role"])) if b["role"] else "")
                    txt1(s, (x, y + 286, w, 36), attr, 17, cc["muted"], 400)
            elif b["type"] == "chart":
                add_native_chart(s, b["box"], b["spec"])
            elif b["type"] == "table":
                x, y, w, h = b["box"]; cc = card_colors(theme); cols = b.get("columns", []); rows = b.get("rows", [])
                ncol, nrow = len(cols), len(rows) + 1
                if ncol and rows:
                    gt = s.shapes.add_table(nrow, ncol, Emu(emu(x)), Emu(emu(y)), Emu(emu(w)), Emu(emu(min(h, 58 + len(rows) * 52)))).table
                    gt.first_row = False; gt.horz_banding = False
                    for ci in range(ncol):
                        cell = gt.cell(0, ci); cell.text = str(cols[ci])
                        style_cell(cell, cc["heading"], 15, True, cc["accent_wash"], center=(ci > 0))
                    for ri, r in enumerate(rows, 1):
                        c0 = gt.cell(ri, 0); c0.text = str(r.get("label", "")); style_cell(c0, cc["heading"], 15, True, cc["fill"])
                        cells = r.get("cells", [])
                        for ci in range(1, ncol):
                            cv = cells[ci - 1] if ci - 1 < len(cells) else ""
                            mark = {"○": "✓", "◯": "✓", "✔": "✓", "×": "✗", "✕": "✗", "△": "△", "-": "△", "ー": "△"}.get(str(cv))
                            cell = gt.cell(ri, ci); cell.text = mark if mark else str(cv)
                            mcol = cc["up"] if mark == "✓" else cc["down"] if mark == "✗" else cc["muted"] if mark == "△" else cc["body"]
                            style_cell(cell, mcol, 16 if mark else 15, bool(mark), cc["fill"], center=True)
            elif b["type"] == "steps":
                x, y, w, h = b["box"]; cc = card_colors(theme); sts = b["steps"]; n = max(1, len(sts)); cw2 = w / n
                for i, st in enumerate(sts):
                    sx = x + i * cw2
                    txt1(s, (sx, y, cw2 - 16, 52), f"{i+1:02d}", 40, cc["accent"], 800, line=1.0)
                    rl = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(emu(sx)), Emu(emu(y + 58)), Emu(emu(36)), Emu(emu(3)))
                    rl.shadow.inherit = False; rl.line.fill.background(); rl.fill.solid(); rl.fill.fore_color.rgb = rgb(cc["accent"])
                    txt1(s, (sx, y + 74, cw2 - 16, 44), st.get("title", ""), 18, cc["heading"], 700, line=1.35)
                    txt1(s, (sx, y + 118, cw2 - 16, 90), st.get("desc", ""), 15, cc["body"], 400, line=1.5)
            elif b["type"] == "timeline":
                x, y, w, h = b["box"]; cc = card_colors(theme); ms = b["items"]; n = max(1, len(ms)); cw2 = w / n
                if n > 1:
                    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(emu(x + cw2 / 2)), Emu(emu(y + 30)), Emu(emu(w - cw2)), Emu(emu(2)))
                    ln.shadow.inherit = False; ln.line.fill.background(); ln.fill.solid(); ln.fill.fore_color.rgb = rgb(cc["border"])
                for i, m in enumerate(ms):
                    cx = x + i * cw2 + cw2 / 2
                    txt1(s, (x + i * cw2, y, cw2, 26), m.get("date", ""), 15, cc["accent"], 600, line=1.2, align="center")
                    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Emu(emu(cx - 8)), Emu(emu(y + 22)), Emu(emu(16)), Emu(emu(16)))
                    dot.shadow.inherit = False; dot.line.color.rgb = rgb(cc["bg"]); dot.line.width = Pt(3)
                    dot.fill.solid(); dot.fill.fore_color.rgb = rgb(cc["accent"])
                    txt1(s, (x + i * cw2, y + 48, cw2, 60), m.get("label", ""), 16, cc["heading"], 700, line=1.4, align="center")
            elif b["type"] == "matrix":
                x, y, w, h = b["box"]; cc = card_colors(theme); ax = b.get("axis_x") or {}; ay = b.get("axis_y") or {}
                ps = min(h - 34, 470); px0 = x + 138; py0 = y + 4
                pl = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(emu(px0)), Emu(emu(py0)), Emu(emu(ps)), Emu(emu(ps)))
                pl.shadow.inherit = False; pl.fill.solid(); pl.fill.fore_color.rgb = rgb(cc["fill"])
                pl.line.color.rgb = rgb(cc["border"]); pl.line.width = Pt(1)
                for (bx, by, bw, bh) in [(px0 + ps / 2, py0, 1.4, ps), (px0, py0 + ps / 2, ps, 1.4)]:
                    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(emu(bx)), Emu(emu(by)), Emu(emu(bw)), Emu(emu(bh)))
                    ln.shadow.inherit = False; ln.line.fill.background(); ln.fill.solid(); ln.fill.fore_color.rgb = rgb(cc["border"])
                for it in b.get("items", []):
                    ix = px0 + _f01(it.get("x")) * ps; iy = py0 + (1 - _f01(it.get("y"))) * ps
                    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Emu(emu(ix - 9)), Emu(emu(iy - 9)), Emu(emu(18)), Emu(emu(18)))
                    dot.shadow.inherit = False; dot.line.fill.background(); dot.fill.solid(); dot.fill.fore_color.rgb = rgb(cc["accent"])
                    txt1(s, (ix + 14, iy - 12, 230, 26), it.get("label", ""), 15, cc["heading"], 700)
                txt1(s, (px0, py0 + ps + 6, 220, 24), ax.get("low", ""), 13, cc["muted"], 400)
                txt1(s, (px0 + ps - 220, py0 + ps + 6, 220, 24), ax.get("high", ""), 13, cc["muted"], 400, align="center")
                txt1(s, (x, py0 + ps - 18, 126, 22), ay.get("low", ""), 13, cc["muted"], 400)
                txt1(s, (x, py0, 126, 22), ay.get("high", ""), 13, cc["muted"], 400)
            elif b["type"] == "funnel":
                x, y, w, h = b["box"]; cc = card_colors(theme); stg = b.get("stages", []); n = max(1, len(stg))
                pal = chart_palette(theme); sh = (h - 8) / n; topw = w * 0.72; cx = x + w / 2; minw = w * 0.14
                fvals = [_fnum(st.get("value")) for st in stg]; mx = max([v for v in fvals if v] or [0])
                for i, st in enumerate(stg):
                    v = fvals[i]
                    wi = (max(minw, topw * v / mx) if (v is not None and mx > 0)
                          else max(minw, topw * (1 - 0.55 * (i / max(1, n - 1)))))
                    yy = y + i * sh
                    bar = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(emu(cx - wi / 2)), Emu(emu(yy)), Emu(emu(wi)), Emu(emu(sh - 10)))
                    bar.shadow.inherit = False; bar.line.fill.background(); bar.fill.solid(); bar.fill.fore_color.rgb = rgb(pal[i % len(pal)])
                    val = ("　" + _fmt(st.get("value"))) if str(st.get("value", "")).strip() else ""
                    txt1(s, (cx - wi / 2, yy + (sh - 10) / 2 - 16, wi, 32), str(st.get("label", "")) + val, 17, "#ffffff", 700, align="center")
            elif b["type"] == "visual":
                x, y, w, h = b["box"]; cc = card_colors(theme); hh = min(h, 440); nm = b.get("icon"); pts = b.get("points", [])
                pw = w * 0.42; panelx = x + w - pw
                pan = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(emu(panelx)), Emu(emu(y)), Emu(emu(pw)), Emu(emu(hh)))
                pan.shadow.inherit = False; pan.line.fill.background(); pan.fill.solid(); pan.fill.fore_color.rgb = rgb(cc["accent_wash"])
                png = icon_png.get((nm, accent, 128)) if nm else None
                if png:
                    s.shapes.add_picture(str(png), Emu(emu(panelx + pw / 2 - 64)), Emu(emu(y + hh / 2 - 64)), Emu(emu(128)), Emu(emu(128)))
                yy = y + max(0.0, (hh - len(pts) * 46) / 2)
                for pnt in pts:
                    mk = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(emu(x)), Emu(emu(yy + 9)), Emu(emu(8)), Emu(emu(8)))
                    mk.shadow.inherit = False; mk.line.fill.background(); mk.fill.solid(); mk.fill.fore_color.rgb = rgb(cc["accent"])
                    txt1(s, (x + 18, yy, w - pw - 44, 44), pnt, 19, cc["body"], 400, line=1.5); yy += 46
    prs.save(str(out_path))

    # フォント埋め込み(subset を family でまとめる)
    by_family = {}
    for (fam, bold), data in subsets.items():
        by_family.setdefault(fam, {})["bold" if bold else "regular"] = data
    fonts.embed_fonts_into_pptx(out_path, by_family)


# ============================================================
# API 向け 再利用関数
# ============================================================
def render_deck(deck, theme, outdir):
    """deck+theme → 各スライドPNG。戻り: {"pngs":[paths], "subsets": {...}}(pptx出力で再利用)"""
    outdir = pathlib.Path(outdir); outdir.mkdir(parents=True, exist_ok=True)
    # 旧PNGを掃除(枚数を減らす編集後に stale ページが残る/順序が崩れるのを防ぐ)
    for old in outdir.glob("s*.png"):
        old.unlink()
    subsets = build_font_assets(deck, theme)
    font_css = font_face_css(subsets)
    pngs = render_pngs(deck, theme, outdir, font_css)
    return {"pngs": pngs, "subsets": subsets}


def export_pdf(png_paths, out_path):
    from PIL import Image
    Image.init()  # PDF保存はJPEGエンコーダ登録が必要(遅延登録されず KeyError:'JPEG' になる対策)
    if not png_paths:
        raise ValueError("PDF化するページがありません")
    imgs = [Image.open(p).convert("RGB") for p in png_paths]
    imgs[0].save(str(out_path), save_all=True, append_images=imgs[1:], resolution=144.0)
    for im in imgs:
        im.close()
    return out_path


def _display(p: pathlib.Path) -> str:
    """表示用のパス文字列。 ROOT.parent 配下なら相対、 外なら絶対のまま返す。

    --out に相対パス、 または ROOT.parent の外にある絶対パスを渡すと
    relative_to() が ValueError で落ちる。 表示のためだけに処理を止めないようにする。
    (実測: 相対 --out を渡すと PNG 生成後・PPTX 書き出しの手前で中断していた)
    """
    resolved = p.resolve()
    try:
        return str(resolved.relative_to(ROOT.parent))
    except ValueError:
        return str(resolved)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--deck", default=str(ROOT / "samples" / "sample_deck.json"))
    ap.add_argument("--out",  default=str(ROOT / "out"))
    args = ap.parse_args()

    # --deck は人が用意する。BOM 付きで保存されていても読めるよう utf-8-sig で読む
    # (BOM が無ければ utf-8 と同じ挙動)。
    deck = json.loads(pathlib.Path(args.deck).read_text(encoding="utf-8-sig"))
    theme = yaml.safe_load((ROOT / "themes" / f"{deck['deck']['theme']}.yaml").read_text(encoding="utf-8"))
    outdir = pathlib.Path(args.out); outdir.mkdir(parents=True, exist_ok=True)

    print(f"[deck] {deck['deck']['title']}  theme={deck['deck']['theme']}  slides={len(deck['slides'])}")
    subsets = build_font_assets(deck, theme)
    print("[fonts] subset: " + ", ".join(f"{fam}{'/B' if b else ''}={len(d)//1024}KB" for (fam, b), d in subsets.items()))
    font_css = font_face_css(subsets)

    for pth in render_pngs(deck, theme, outdir, font_css):
        print(f"  PNG  {_display(pth)}  ({pth.stat().st_size} bytes)")
    pptx_path = outdir / "deck.pptx"
    render_pptx(deck, theme, pptx_path, subsets)
    print(f"  PPTX {_display(pptx_path)}  ({pptx_path.stat().st_size} bytes)  ← フォント埋め込み・編集可")
    print("[done]")


if __name__ == "__main__":
    sys.exit(main())
