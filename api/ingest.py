"""ソース取り込み: URL / PDF / PPTX / DOCX / テキスト → プレーンテキスト抽出。

抽出は課金ゼロ(ローカル解析)。スライド本文はこの抽出テキストのみを根拠にする(捏造防止)。
"""
from __future__ import annotations
import html as _html
import pathlib
import re

MAX_CHARS = 60000   # 忠実プリセット(全量保持)で長文ソースが途中で切れないよう拡大

# テキストとして妥当かの判定に使う制御文字 (タブ・改行・復帰は除く)。
# これが混じっている = その文字コードでの解釈が誤り、とみなす。
_CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

# Markdown の見出し記号 (行頭の # 〜 ######)。 インデントは保つ。
_MD_HEADING_MARK = re.compile(r"^(\s{0,3})#{1,6}[ \t]+", re.M)


def _clip(t: str) -> str:
    t = re.sub(r"[ \t]+", " ", t or "")
    t = re.sub(r"\n{3,}", "\n\n", t).strip()
    return t[:MAX_CHARS]


def from_text(text: str) -> str:
    return _clip(text)


def _read_text_file(path: pathlib.Path) -> str:
    """テキストファイルを読む。 判定できなければ黙って壊さず、 理由を出して止める。

    以前は errors="ignore" で読んでいた。 cp932 で保存されたファイルを渡すと
    不正バイトが捨てられ、 「日本語」が「{」になるなど中身がほぼ消えたまま
    処理だけ進んでいた (実測)。 抽出テキストはスライド本文の唯一の
    根拠なので、 ここで欠落すると中身の無い成果物が黙って出来上がる。

    utf-8-sig は BOM 付きも扱える。 日本語 Windows の実情に合わせて cp932 も試す。

    ただし cp932 は不正なバイト列でもほとんど例外を出さずに解釈してしまう
    (実測: バイナリを渡すと文字化けした文字列が返る)。 UnicodeDecodeError
    だけを判定に使うと、 バイナリを渡した時に文字化けが通ってしまうため、
    制御文字が混じっていたら「そのコードでの解釈は誤り」とみなす。
    """
    for enc in ("utf-8-sig", "cp932"):
        try:
            text = path.read_text(encoding=enc)
        except UnicodeDecodeError:
            continue
        if _CONTROL_CHARS.search(text):
            continue
        return text
    raise ValueError(
        f"文字コードを判定できませんでした: {path.name}\n"
        "UTF-8 または Shift_JIS(cp932) で保存し直してください。"
    )


def from_url(url: str) -> str:
    """利用者が指定した URL の本文を取り込む。自分が権利を持つ、または利用許諾のあるページに限ること
    （第三者の記事は要約と出典明記の範囲で扱う）。"""
    import httpx
    r = httpx.get(url, follow_redirects=True, timeout=30,
                  headers={"User-Agent": "slide-studio/1.0"})
    r.raise_for_status()
    doc = r.text
    doc = re.sub(r"(?is)<(script|style|noscript|nav|footer|header)[^>]*>.*?</\1>", " ", doc)
    doc = re.sub(r"(?is)<br\s*/?>", "\n", doc)
    doc = re.sub(r"(?is)</(p|div|li|h[1-6]|tr)>", "\n", doc)
    doc = re.sub(r"(?is)<[^>]+>", " ", doc)
    return _clip(_html.unescape(doc))


def from_file(path: pathlib.Path) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        import fitz
        d = fitz.open(str(path))
        return _clip("\n".join(pg.get_text() for pg in d))
    if ext in (".pptx", ".ppt"):
        from pptx import Presentation
        prs = Presentation(str(path))
        parts = []
        for s in prs.slides:
            for sh in s.shapes:
                if sh.has_text_frame:
                    parts.append(sh.text_frame.text)
        return _clip("\n".join(parts))
    if ext == ".docx":
        import docx
        return _clip("\n".join(p.text for p in docx.Document(str(path)).paragraphs))
    if ext in (".md", ".markdown"):
        # 見出し記号だけ落とす。 文言は残す (抽出テキストは本文の唯一の根拠なので
        # 情報は減らさない)。 `#` がそのままスライドに描画されるのを防ぐだけ。
        return _clip(_MD_HEADING_MARK.sub(r"\1", _read_text_file(path)))
    # txt / その他テキスト
    return _clip(_read_text_file(path))
